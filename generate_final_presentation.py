import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path="PhantomNet_Final_Presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette Definitions (Cybersecurity Academic Theme)
    NAVY_DARK = RGBColor(0x0F, 0x1E, 0x36)      # Deep Navy Canvas
    NAVY_PRIMARY = RGBColor(0x1B, 0x3A, 0x6B)   # Primary Navy Headers
    BLUE_ACCENT = RGBColor(0x00, 0x66, 0xCC)    # Sharp Cobalt Accent
    CYAN_TECH = RGBColor(0x02, 0x84, 0xC7)      # Cyan Tech Tracker
    BG_LIGHT = RGBColor(0xF8, 0xFA, 0xFC)       # Clean Off-white Slate
    CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)        # White Card Container
    CARD_BORDER = RGBColor(0xCB, 0xD5, 0xE1)    # Crisp Subtle Border
    TEXT_MAIN = RGBColor(0x0F, 0x17, 0x2A)      # Main High-contrast Slate Text
    TEXT_MUTED = RGBColor(0x47, 0x55, 0x69)     # Secondary Slate Text
    TEXT_LIGHT = RGBColor(0x94, 0xA3, 0xB8)     # Metadata Slate Light
    GREEN_SUCCESS = RGBColor(0x15, 0x80, 0x3D)  # Green 700 Verification
    RED_ALERT = RGBColor(0xB9, 0x1C, 0x1C)      # Red 700 Security Alert / Boundary
    AMBER_WARN = RGBColor(0xB4, 0x53, 0x09)     # Amber 700 Warning
    TAG_BG_BLUE = RGBColor(0xEE, 0xF2, 0xF6)
    TAG_BG_GREEN = RGBColor(0xDC, 0xFC, 0xE7)
    TAG_BG_RED = RGBColor(0xFE, 0xE2, 0xE2)

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text, slide_num_str):
        # Category Tracker Pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.0), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Calibri"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN_TECH

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(10.5), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Calibri"
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY_PRIMARY

        # Slide Number
        num_box = slide.shapes.add_textbox(Inches(11.8), Inches(0.5), Inches(0.8), Inches(0.4))
        tf_num = num_box.text_frame
        tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = slide_num_str
        p_num.font.name = "Calibri"
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = TEXT_LIGHT

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def add_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    print("Generating comprehensive, audited 18-slide deck...")

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Executive Theme)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, NAVY_DARK)

    badge = add_card(slide1, Inches(0.8), Inches(0.8), Inches(4.5), Inches(0.4), bg_color=RGBColor(0x1E, 0x2E, 0x48), border_color=CYAN_TECH)
    tf = badge.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "FINAL YEAR PROJECT PRESENTATION  |  2026"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_TECH

    title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(1.3))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PHANTOMNET"
    p.font.name = "Calibri"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sub_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(0.6))
    tf = sub_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "AI-Assisted Distributed Honeypot & Threat Analysis Platform"
    p.font.name = "Calibri"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = CYAN_TECH

    desc_box = slide1.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(11.5), Inches(0.8))
    tf = desc_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "An integrated defensive cybersecurity engineering platform converting adversarial interactions into real-time telemetry, machine-learning threat classification, host-level active defense, and automated incident response playbooks."
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    stages = [
        ("1. DECEIVE", "4 Custom Micro-Honeypots\nSSH, HTTP, FTP, SMTP"),
        ("2. INGEST", "Resilient Spooling &\nSHA-256 HMAC Dispatch"),
        ("3. DETECT", "12D Flow Features &\nRandom Forest Ensemble"),
        ("4. RESPOND", "0.8 Safety Firewall Block &\nSentinel Sigma/STIX Rules")
    ]
    for i, (stitle, sdesc) in enumerate(stages):
        x = Inches(0.8 + i * 2.95)
        card = add_card(slide1, x, Inches(4.2), Inches(2.75), Inches(1.2), bg_color=RGBColor(0x16, 0x24, 0x3C), border_color=RGBColor(0x2A, 0x43, 0x68))
        tf = card.text_frame
        tf.margin_top = Inches(0.12)
        tf.margin_left = Inches(0.15)
        p1 = tf.paragraphs[0]
        p1.text = stitle
        p1.font.name = "Calibri"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = CYAN_TECH
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.name = "Calibri"
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

    team_box = slide1.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.7), Inches(0.9))
    tf = team_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PROJECT TEAM & ROLES:"
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_TECH
    p2 = tf.add_paragraph()
    p2.text = "Kasukurthi Sriram (Lead Architect)  •  Muramreddy Vivekananda Reddy (Infrastructure)  •  Nattala Vikranth Chakravarthi (ML Engineer)  •  Satti Sai Ram Manideep Reddy (Frontend Lead)"
    p2.font.name = "Calibri"
    p2.font.size = Pt(12)
    p2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    add_notes(slide1, "Speaker Introduction (30 seconds):\n'Good morning esteemed examiners. Today we present PhantomNet: an AI-Assisted Distributed Honeypot and Threat Analysis Platform. Traditional security defenses are predominantly reactive, identifying adversaries only after perimeter compromise. PhantomNet implements an integrated defensive engineering pipeline: attracting adversaries via four custom multi-protocol micro-honeypots, capturing cryptographically authenticated telemetry, classifying threats via an in-domain machine learning ensemble with a strict zero-false-positive safety gate, and automatically synthesizing MITRE ATT&CK mapped incident playbooks and SIEM detection rules.'")

    # =========================================================================
    # SLIDE 2: Problem Statement — The Reactive Security Deficit
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, BG_LIGHT)
    add_header(slide2, "The Problem — The Reactive Security Deficit", "Problem Statement", "02 / 18")

    stat_items = [
        ("207 Days", "Average Breach Dwell Time", "Adversaries perform extensive reconnaissance and lateral movement before detection (IBM Cost of a Data Breach Report).", RED_ALERT),
        ("Alert Fatigue", "Thousands of Daily Logs", "SOC analysts are inundated with perimeter alerts lacking behavioral context and adversary attribution.", AMBER_WARN),
        ("Asymmetric Cost", "$100K+ Commercial Tools", "Enterprise deception technology is inaccessible to SMEs, educational labs, and public sector networks.", NAVY_PRIMARY)
    ]
    for i, (val, sub, desc, color) in enumerate(stat_items):
        c = add_card(slide2, Inches(0.8 + i * 3.95), Inches(1.5), Inches(3.75), Inches(1.8))
        tf = c.text_frame
        tf.margin_left = tf.margin_top = Inches(0.2)
        p1 = tf.paragraphs[0]
        p1.text = val
        p1.font.name = "Calibri"
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = "Calibri"
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_MAIN
        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = "Calibri"
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED

    bottom_cards = [
        ("🚨 Detection Occurs After Damage", "Standard firewalls and IDSes trigger alerts when malicious payloads have already penetrated boundaries. They offer zero safe sandbox interaction to observe attacker tactics without risking assets."),
        ("🔍 Zero Attacker Context & Tooling Insight", "Signature-based rules flag known IP/hash indicators but fail against zero-day reconnaissance. They reveal nothing about attacker tools, credential lists, and active command execution sequences."),
        ("📉 The Open-Source Disconnect", "Freely available honeypots (e.g., standalone Cowrie/Dionaea) capture raw logs into text files but lack unified streaming telemetry, machine learning risk classification, automated defense, and SOC integration.")
    ]
    for i, (head, body) in enumerate(bottom_cards):
        c = add_card(slide2, Inches(0.8 + i * 3.95), Inches(3.55), Inches(3.75), Inches(3.2))
        tf = c.text_frame
        tf.margin_left = tf.margin_top = Inches(0.2)
        p1 = tf.paragraphs[0]
        p1.text = head
        p1.font.name = "Calibri"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = NAVY_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    add_notes(slide2, "Examiner Key Point:\n'The core problem is not a lack of security logs; it is a lack of high-fidelity, early-stage adversary context. By the time perimeter alarms sound, the dwell time clock has been ticking for months. Deception provides an unfair defensive advantage by converting untrusted network space into tripwires, but existing tools leave honeypots as isolated data silos.'")

    # =========================================================================
    # SLIDE 3: Proposed Solution — Integrated Defensive Architecture
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, BG_LIGHT)
    add_header(slide3, "The Solution — Integrated Defensive Cybersecurity Architecture", "Solution Overview", "03 / 18")

    stmt_card = add_card(slide3, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.9), bg_color=TAG_BG_BLUE, border_color=CYAN_TECH)
    tf = stmt_card.text_frame
    tf.margin_left = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "PhantomNet combines controlled network deception with resilient streaming ingestion, an explainable machine learning scoring engine, host-level active containment, and automated detection rule engineering."
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    pillars = [
        ("01. Controlled Deception Mesh", "4 Custom Micro-Honeypots", "• Emulates SSH (2222), HTTP (8080), FTP (2121), and SMTP (2525)\n• Custom lightweight Python architecture\n• High-fidelity service interaction traps\n• Hardened container isolation with cap-drop", BLUE_ACCENT),
        ("02. Resilient Telemetry Pipeline", "HMAC Signatures & Local Spooling", "• Cryptographic SHA-256 HMAC event signatures\n• Local disk spooling prevents partition data loss\n• High-throughput FastAPI ingestion endpoint\n• Redis 7 Streams with PEL recovery & DLQ", CYAN_TECH),
        ("03. In-Domain ML Threat Scoring", "12D Flow Features & Dual Threshold", "• Canonical 12-dimensional flow representation\n• Random Forest (97.88% in-domain accuracy)\n• Calibrated Isolation Forest anomaly scoring\n• Strict 0.8 safety threshold: 0.0% false positives", GREEN_SUCCESS),
        ("04. Sentinel Detection Engineering", "MITRE ATT&CK, Sigma, Snort & STIX", "• Automated mapping to 12+ ATT&CK techniques\n• Vendor-agnostic Sigma YAML SIEM rules\n• Network IDS Snort rules with auto SIDs\n• Dual-path generation: Ollama LLM + Jinja2 fallback", NAVY_PRIMARY)
    ]
    for i, (pnum, psub, pbody, pcolor) in enumerate(pillars):
        c = add_card(slide3, Inches(0.8 + i * 2.95), Inches(2.6), Inches(2.8), Inches(4.2))
        tf = c.text_frame
        tf.margin_left = tf.margin_top = Inches(0.2)
        p1 = tf.paragraphs[0]
        p1.text = pnum
        p1.font.name = "Calibri"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = pcolor
        p2 = tf.add_paragraph()
        p2.text = psub
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_MAIN
        p3 = tf.add_paragraph()
        p3.text = pbody
        p3.font.name = "Calibri"
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED

    add_notes(slide3, "Examiner Key Point:\n'Notice that PhantomNet is NOT just a honeypot, and NOT just an ML classifier. It is an end-to-end defensive engineering architecture linking honeypot deception directly to SOC detection engineering and active defense.'")

    # =========================================================================
    # SLIDE 4: System Architecture — The Verified 6-Layer Pipeline
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, BG_LIGHT)
    add_header(slide4, "System Architecture — Verified 6-Layer Pipeline", "Architecture", "04 / 18")

    layers = [
        ("LAYER 1: DECEPTION LAYER", "4 Custom Micro-Honeypots: SSH (2222), HTTP (8080), FTP (2121), SMTP (2525) running in isolated bridge networks.", NAVY_PRIMARY),
        ("LAYER 2: INGESTION LAYER", "Local Spooler (`spooler.py`) → HMAC-SHA256 Signature → FastAPI Gateway (`/api/v1/ingest/event`).", BLUE_ACCENT),
        ("LAYER 3: STREAM & PERSISTENCE", "PostgreSQL 15 (Relational events & playbooks) + Redis 7 Streams (Consumer groups, PEL recovery, DLQ).", CYAN_TECH),
        ("LAYER 4: ML & ANALYTICS LAYER", "12D Flow Extractor (`feature_extractor.py`) → Random Forest Classifier + Calibrated Isolation Forest.", GREEN_SUCCESS),
        ("LAYER 5: RESPONSE & SENTINEL LAYER", "0.8 Confidence Threshold → Host Firewall (`iptables`/`netsh`) + MITRE ATT&CK Mapping + Sigma/Snort Rule Synthesis.", AMBER_WARN),
        ("LAYER 6: SOC OBSERVABILITY LAYER", "React 19 Dashboard (11 Production Views, Vite/Rolldown build) + Real-Time WebSocket Streaming (`/ws`).", NAVY_DARK)
    ]
    for i, (ltitle, ldesc, lcolor) in enumerate(layers):
        y = Inches(1.5 + i * 0.9)
        card = add_card(slide4, Inches(0.8), y, Inches(11.7), Inches(0.78), bg_color=CARD_BG, border_color=lcolor)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.1)
        p1 = tf.paragraphs[0]
        p1.text = ltitle
        p1.font.name = "Calibri"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = lcolor
        p2 = tf.add_paragraph()
        p2.text = ldesc
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN

    add_notes(slide4, "Examiner Key Point:\n'Every layer in this 6-layer pipeline is verified in the active codebase: 268 backend Python files, 160 route endpoints across 27 routers, and 11 production views in the frontend. Notice the unidirectional data flow: honeypots cannot read the database or access the internal broker.'")

    # =========================================================================
    # SLIDE 5: Multi-Protocol Deception Layer — Custom Micro-Honeypots
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, BG_LIGHT)
    add_header(slide5, "Multi-Protocol Deception Layer — Custom Micro-Honeypots", "Deception Layer", "05 / 18")

    corr_banner = add_card(slide5, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.65), bg_color=TAG_BG_BLUE, border_color=CYAN_TECH)
    tf = corr_banner.text_frame
    tf.margin_left = Inches(0.2)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "CORRECTION & ARCHITECTURAL CLARIFICATION: PhantomNet utilizes four bespoke lightweight Python micro-honeypots engineered specifically for controlled deception and telemetry capture — NOT third-party Cowrie/Dionaea daemons."
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    h_cards = [
        ("SSH Honeypot", "Port 2222 • AsyncSSH", "• Simulated Unix filesystem\n• Realistic login banners\n• Credential harvesting\n• Regex credential sanitization\n• Command history recording", "backend/honeypots/ssh/"),
        ("HTTP Honeypot", "Port 8080 • Custom Server", "• Fake admin & wp-admin logins\n• Form submission traps\n• SQL injection tripwires\n• Path traversal detection\n• User-agent & header inspection", "backend/honeypots/http/"),
        ("FTP Honeypot", "Port 2121 • pyftpdlib", "• Emulates vsFTPd 3.0.3 banner\n• Simulated file directories\n• Anonymous access monitoring\n• USER / PASS brute-force logging\n• Session duration tracking", "backend/honeypots/ftp/"),
        ("SMTP Honeypot", "Port 2525 • aiosmtpd", "• Postfix / Exim emulation\n• Fake mail relay traps\n• Mass-mailing probe logging\n• EICAR malware test strings\n• Full conversation telemetry", "backend/honeypots/smtp/")
    ]
    for i, (htitle, hsub, hbody, hpath) in enumerate(h_cards):
        c = add_card(slide5, Inches(0.8 + i * 2.95), Inches(2.3), Inches(2.8), Inches(3.5))
        tf = c.text_frame
        tf.margin_left = tf.margin_top = Inches(0.18)
        p1 = tf.paragraphs[0]
        p1.text = htitle
        p1.font.name = "Calibri"
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = NAVY_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = hsub
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = BLUE_ACCENT
        p3 = tf.add_paragraph()
        p3.text = hbody
        p3.font.name = "Calibri"
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED
        p4 = tf.add_paragraph()
        p4.text = f"\nPath: {hpath}"
        p4.font.name = "Calibri"
        p4.font.size = Pt(9)
        p4.font.color.rgb = TEXT_LIGHT

    res_card = add_card(slide5, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    tf = res_card.text_frame
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.1)
    p1 = tf.paragraphs[0]
    p1.text = "TELEMETRY RESILIENCE & INTEGRITY:"
    p1.font.name = "Calibri"
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = GREEN_SUCCESS
    p2 = tf.add_paragraph()
    p2.text = "• Local Disk Spooling (`spooler.py`): Events are spooled locally if the backend is unreachable, preventing data loss during network partitions.\n• Cryptographic Authenticity (`event_dispatcher.py`): Ingestion payloads are signed via SHA-256 HMAC headers, preventing rogue injection."
    p2.font.name = "Calibri"
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MAIN

    add_notes(slide5, "Viva Defense Question Anticipation:\n'Examiner: Why build custom honeypots instead of running Cowrie or Dionaea?'\n'Answer: Standard daemons like Cowrie require heavy runtime dependencies and external logging forwarders. Our custom micro-honeypots integrate native SHA-256 HMAC event signing and offline disk spooling directly into the connection handler, ensuring tamper-proof telemetry and resilience against network partition.'")

    # =========================================================================
    # SLIDE 6: End-to-End Attack & Detection Flow
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, BG_LIGHT)
    add_header(slide6, "End-to-End Attack & Detection Lifecycle", "Event Pipeline", "06 / 18")

    flow_steps = [
        ("1. Adversary Probe", "Port scan, SSH brute force, or SQL injection hits exposed honeypot service."),
        ("2. Honeypot Trap", "Service interacts, sanitizes credentials, signs event with SHA-256 HMAC (`spooler.py`)."),
        ("3. Gateway Ingest", "FastAPI (`/api/v1/ingest/event`) verifies HMAC, commits to PostgreSQL, pushes to Redis Stream."),
        ("4. Flow Extraction", "Event consumer extracts canonical 12D flow vector with rolling window statistics."),
        ("5. ML Inference", "Random Forest + Calibrated Isolation Forest compute threat confidence score (0.0 – 1.0)."),
        ("6. Safety Gating", "Score evaluated against dual thresholds: 0.5 (Log/Alert) vs 0.8 (Automated Defense Trigger)."),
        ("7. Active Mitigation", "Host firewall command (`iptables` / `netsh`) dispatched to isolate attacking IP address."),
        ("8. Sentinel Playbook", "Incident mapped to MITRE ATT&CK; Sigma YAML, Snort rules, and STIX 2.1 bundles generated."),
        ("9. SOC Dashboard", "Live event stream, threat score gauge, and playbook rendered via WebSocket on React 19 UI.")
    ]
    for i, (step_title, step_desc) in enumerate(flow_steps):
        row = i // 3
        col = i % 3
        x = Inches(0.8 + col * 3.95)
        y = Inches(1.5 + row * 1.8)
        c = add_card(slide6, x, y, Inches(3.75), Inches(1.6))
        tf = c.text_frame
        tf.margin_left = tf.margin_top = Inches(0.18)
        p1 = tf.paragraphs[0]
        p1.text = step_title
        p1.font.name = "Calibri"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = BLUE_ACCENT if i < 3 else (GREEN_SUCCESS if i < 6 else NAVY_PRIMARY)
        p2 = tf.add_paragraph()
        p2.text = step_desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

    add_notes(slide6, "Examiner Key Point:\n'This 9-step lifecycle illustrates how a raw TCP packet from an adversary is deterministically transformed into a classified threat event, an active firewall rule, an analyst playbook, and an exportable STIX 2.1 intelligence bundle.'")

    # =========================================================================
    # SLIDE 7: ML Detection Engine — Canonical Representation & Strategy
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, BG_LIGHT)
    add_header(slide7, "ML Detection Engine — Canonical Representation & Strategy", "Machine Learning", "07 / 18")

    left_c = add_card(slide7, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf = left_c.text_frame
    tf.margin_left = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "CANONICAL 12-DIMENSIONAL FLOW FEATURES"
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p_sub = tf.add_paragraph()
    p_sub.text = "Implemented in `backend/ml/feature_extractor.py`:\n"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(11)
    p_sub.font.bold = True
    p_sub.font.color.rgb = CYAN_TECH

    features = [
        "1. Destination Port Normalization (Service categorization)",
        "2. Connection Frequency (Flows within rolling window)",
        "3. Flow Duration (Connection persistence in seconds)",
        "4. Total Transferred Bytes (Volume metrics)",
        "5. Average Packet Size (Payload density)",
        "6. Protocol One-Hot Encoding (TCP / UDP / ICMP)",
        "7. Rolling Window Packet Rate (Packets per second)",
        "8. Byte Rate (Bytes per second throughput)",
        "9. TCP SYN / ACK Flag Ratio (Handshake behavior)",
        "10. Distinct Port Count (Port scanning indicator)",
        "11. Inter-Arrival Time Variance (Traffic burstiness)",
        "12. Payload Entropy (Information randomness; 0 if empty)"
    ]
    for feat in features:
        pf = tf.add_paragraph()
        pf.text = feat
        pf.font.name = "Calibri"
        pf.font.size = Pt(10)
        pf.font.color.rgb = TEXT_MAIN

    p_quar = tf.add_paragraph()
    p_quar.text = "\n⚠️ Technical Integrity Note:\nAn earlier 32-feature extractor (`feature_engineering_complete.py`) was audited, deprecated, and quarantined due to target leakage. Only the canonical 12D extractor is used for inference."
    p_quar.font.name = "Calibri"
    p_quar.font.size = Pt(9.5)
    p_quar.font.color.rgb = AMBER_WARN

    right_c = add_card(slide7, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf_r = right_c.text_frame
    tf_r.margin_left = tf_r.margin_top = Inches(0.25)
    pr = tf_r.paragraphs[0]
    pr.text = "ENSEMBLE MODEL ARCHITECTURE & DECISION LOGIC"
    pr.font.name = "Calibri"
    pr.font.size = Pt(14)
    pr.font.bold = True
    pr.font.color.rgb = NAVY_PRIMARY

    models_info = [
        ("Random Forest Classifier (`rf_model.pkl`)", "Supervised classification trained on in-domain honeypot interaction flows. Excels at detecting known reconnaissance, brute force, and exploit patterns with high speed."),
        ("Calibrated Isolation Forest (`if_model.pkl`)", "Unsupervised anomaly detection trained to flag anomalous traffic vectors that deviate from baseline network distributions. Provides calibrated anomaly probability."),
        ("Heuristic Rule Fallback", "Deterministic safety fallback ensuring known threat patterns (e.g., rapid port cycling, credential bursts) are flagged even during ML initialization."),
        ("Runtime Decision Strategy", "Confidence score = Weighted vote combining Random Forest prediction probability and calibrated Isolation Forest score."),
        ("Runtime Path Honesty Note", "LSTM sequence modeling is deactivated / mock in active flow inference. PhantomNet does not claim active LSTM in runtime.")
    ]
    for title, desc in models_info:
        p_t = tf_r.add_paragraph()
        p_t.text = f"\n• {title}"
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = BLUE_ACCENT if "Honesty" not in title else AMBER_WARN
        p_d = tf_r.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MUTED

    add_notes(slide7, "Examiner Key Point:\n'Examiners appreciate academic honesty. When asked about features, state clearly: We use a 12-dimensional canonical flow representation. We identified target leakage in an older 32-feature experimental extractor and actively deprecated it. Furthermore, our runtime ensemble relies on Random Forest and calibrated Isolation Forest; we do not claim active LSTM sequence inference.'")

    # =========================================================================
    # SLIDE 8: ML Performance — Verified In-Distribution & Held-Out Results
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, BG_LIGHT)
    add_header(slide8, "ML Performance — In-Distribution & Safety Gate Evaluation", "Evaluation & Metrics", "08 / 18")

    c_left = add_card(slide8, Inches(0.8), Inches(1.5), Inches(5.7), Inches(4.3))
    tf = c_left.text_frame
    tf.margin_left = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "SYNTHETIC SIMULATOR — 5,000 FLOWS"
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p_sub = tf.add_paragraph()
    p_sub.text = "Evaluated on full in-distribution traffic dataset:\n"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = TEXT_MUTED

    metrics_full = [
        ("Random Forest @ 0.5", "97.88%", "97.87%", "95.00%", "96.41%", "0.89%", "0.9968"),
        ("Isolation Forest", "73.22%", "67.69%", "20.53%", "31.51%", "4.20%", "0.6493"),
        ("Ensemble @ 0.5", "97.36%", "98.37%", "92.73%", "95.47%", "0.66%", "0.9966"),
        ("Ensemble @ 0.8 (BLOCK)", "89.66%", "100.0%", "65.53%", "79.18%", "0.00%", "0.9966")
    ]
    for model_name, acc, prec, rec, f1, fpr, auc in metrics_full:
        pm = tf.add_paragraph()
        is_block = "0.8" in model_name
        pm.text = f"• {model_name}:\n   Acc: {acc} | Prec: {prec} | Rec: {rec} | F1: {f1}\n   FPR: {fpr} | AUC-ROC: {auc}"
        pm.font.name = "Calibri"
        pm.font.size = Pt(10)
        pm.font.bold = is_block
        pm.font.color.rgb = GREEN_SUCCESS if is_block else TEXT_MAIN

    c_right = add_card(slide8, Inches(6.8), Inches(1.5), Inches(5.7), Inches(4.3))
    tf_r = c_right.text_frame
    tf_r.margin_left = tf_r.margin_top = Inches(0.2)
    pr = tf_r.paragraphs[0]
    pr.text = "HELD-OUT UNSEEN SPLIT — 1,000 FLOWS (20%, Seed 42)"
    pr.font.name = "Calibri"
    pr.font.size = Pt(13)
    pr.font.bold = True
    pr.font.color.rgb = NAVY_PRIMARY

    pr_sub = tf_r.add_paragraph()
    pr_sub.text = "Evaluated on strictly held-out unseen test samples:\n"
    pr_sub.font.name = "Calibri"
    pr_sub.font.size = Pt(10)
    pr_sub.font.color.rgb = TEXT_MUTED

    metrics_held = [
        ("Random Forest @ 0.5", "93.00%", "94.23%", "81.67%", "87.50%", "2.14%", "0.9793"),
        ("Isolation Forest", "73.40%", "71.25%", "19.00%", "30.00%", "3.29%", "0.6084"),
        ("Ensemble @ 0.5", "92.70%", "95.22%", "79.67%", "86.75%", "1.71%", "0.9785"),
        ("Ensemble @ 0.8 (BLOCK)", "87.50%", "100.0%", "58.33%", "73.68%", "0.00%", "0.9785")
    ]
    for model_name, acc, prec, rec, f1, fpr, auc in metrics_held:
        pm = tf_r.add_paragraph()
        is_block = "0.8" in model_name
        pm.text = f"• {model_name}:\n   Acc: {acc} | Prec: {prec} | Rec: {rec} | F1: {f1}\n   FPR: {fpr} | AUC-ROC: {auc}"
        pm.font.name = "Calibri"
        pm.font.size = Pt(10)
        pm.font.bold = is_block
        pm.font.color.rgb = GREEN_SUCCESS if is_block else TEXT_MAIN

    banner = add_card(slide8, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.9), bg_color=TAG_BG_GREEN, border_color=GREEN_SUCCESS)
    tf_b = banner.text_frame
    tf_b.margin_left = Inches(0.25)
    tf_b.margin_top = Inches(0.1)
    pb1 = tf_b.paragraphs[0]
    pb1.text = "THE ZERO-FALSE-POSITIVE SAFETY GATE (THRESHOLD = 0.8):"
    pb1.font.name = "Calibri"
    pb1.font.size = Pt(11)
    pb1.font.bold = True
    pb1.font.color.rgb = GREEN_SUCCESS
    pb2 = tf_b.add_paragraph()
    pb2.text = "At the standard 0.5 threshold, the ensemble achieves high sensitivity (Recall 92.7%, F1 95.5%). For automated mitigation (firewall IP blocking), PhantomNet enforces a strict 0.8 confidence threshold. On held-out unseen data, this achieves 100.0% Precision and 0.00% FPR (0 false positives), eliminating denial-of-service risk to benign traffic."
    pb2.font.name = "Calibri"
    pb2.font.size = Pt(10.5)
    pb2.font.color.rgb = TEXT_MAIN

    add_notes(slide8, "Examiner Key Point:\n'Emphasize the dual-threshold strategy. A security tool cannot use a standard 0.5 threshold for automated firewall blocking because even a 1% false positive rate would disconnect legitimate internal servers. By raising the blocking threshold to 0.8, we trade off recall (58.3%) to achieve 100% precision and exactly 0.0% false positives on our held-out test set.'")

    # =========================================================================
    # SLIDE 9: Critical ML Limitation — Cross-Dataset Generalization Boundary
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, BG_LIGHT)
    add_header(slide9, "Critical ML Limitation — Cross-Dataset Generalization", "Limitations & Validity", "09 / 18")

    alert_c = add_card(slide9, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.8), bg_color=TAG_BG_RED, border_color=RED_ALERT)
    tf = alert_c.text_frame
    tf.margin_left = Inches(0.2)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "WHAT THE MODEL DOES — AND DOES NOT — GENERALIZE TO:\nThe current ML models are specialized for the PhantomNet synthetic/honeypot traffic domain. Zero-shot evaluation against heterogeneous public network datasets demonstrates significant performance degradation."
    p.font.name = "Calibri"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = RED_ALERT

    # Embed real heatmap figure on right side if available!
    heatmap_path = "experiments/results/ml9_figures/fig2_cross_domain_transfer_heatmap.png"
    has_heatmap = os.path.exists(heatmap_path)

    left_w = Inches(7.0) if has_heatmap else Inches(11.7)
    c_data = add_card(slide9, Inches(0.8), Inches(2.45), left_w, Inches(2.7))
    tf_d = c_data.text_frame
    tf_d.margin_left = tf_d.margin_top = Inches(0.2)
    pd_title = tf_d.paragraphs[0]
    pd_title.text = "CROSS-DATASET ZERO-SHOT RE-EVALUATION (5,000 Flows Each):"
    pd_title.font.name = "Calibri"
    pd_title.font.size = Pt(12)
    pd_title.font.bold = True
    pd_title.font.color.rgb = NAVY_PRIMARY

    datasets_lines = [
        "• CIC-IDS2017: 28.56% Accuracy | 5.60% F1-Score | 62.23% FPR | 0.2994 AUC-ROC (TP: 106, FP: 2,178)",
        "• UNSW-NB15:   45.20% Accuracy | 40.38% F1-Score | 61.94% FPR | 0.3870 AUC-ROC (TP: 928, FP: 2,168)",
        "• NF-ToN-IoT:  23.58% Accuracy | 32.53% F1-Score | 92.63% FPR | 0.3182 AUC-ROC (TP: 921, FP: 3,242)"
    ]
    for dl in datasets_lines:
        p_dl = tf_d.add_paragraph()
        p_dl.text = dl
        p_dl.font.name = "Calibri"
        p_dl.font.size = Pt(10.5)
        p_dl.font.bold = True
        p_dl.font.color.rgb = RED_ALERT

    p_dnote = tf_d.add_paragraph()
    p_dnote.text = "\nFinding: Models trained on honeypot interactions exhibit high false alarms on bulk enterprise captures due to differing background traffic distributions and packet-size distributions."
    p_dnote.font.name = "Calibri"
    p_dnote.font.size = Pt(10)
    p_dnote.font.color.rgb = TEXT_MUTED

    if has_heatmap:
        slide9.shapes.add_picture(heatmap_path, Inches(8.0), Inches(2.45), Inches(4.5), Inches(2.7))

    # Bottom Academic Context Card
    ctx_card = add_card(slide9, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.6), bg_color=CARD_BG, border_color=CARD_BORDER)
    tf_ctx = ctx_card.text_frame
    tf_ctx.margin_left = Inches(0.25)
    tf_ctx.margin_top = Inches(0.15)
    pc1 = tf_ctx.paragraphs[0]
    pc1.text = "ACADEMIC ROOT CAUSE ANALYSIS & IMPLICATIONS FOR CYBERSECURITY ML:"
    pc1.font.name = "Calibri"
    pc1.font.size = Pt(11)
    pc1.font.bold = True
    pc1.font.color.rgb = NAVY_PRIMARY
    pc2 = tf_ctx.add_paragraph()
    pc2.text = "• Feature Semantic Mismatch: Honeypot interactions exhibit distinct behavioral characteristics compared to bulk enterprise network captures (e.g., higher connection duration variance and distinct port distributions).\n• Need for Domain Adaptation: Rather than claiming false 'universal detection', our findings empirically demonstrate that network intrusion models must undergo domain adaptation and feature re-alignment before cross-domain transfer.\n• Defense Integrity: PhantomNet does NOT claim 98% accuracy on CIC-IDS2017. We report true empirical performance."
    pc2.font.name = "Calibri"
    pc2.font.size = Pt(10.5)
    pc2.font.color.rgb = TEXT_MAIN

    add_notes(slide9, "Viva Defense Masterstroke:\n'When asked about model generalization: Do not get defensive. Present this slide proudly: 'Many academic papers claim 99% accuracy on CIC-IDS2017 by evaluating only on random train-test splits from the exact same dataset. We conducted rigorous cross-dataset zero-shot evaluation and found that models trained on honeypot flow distributions collapse on external PCAPs. This proves that domain shift is a critical boundary, and domain adaptation is necessary for real-world deployment.' This demonstrates true research maturity.'")

    # =========================================================================
    # SLIDE 10: Active Defense & Host Firewall Enforcement
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, BG_LIGHT)
    add_header(slide10, "Active Defense & Host Firewall Enforcement", "Defensive Response", "10 / 18")

    left_c = add_card(slide10, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf = left_c.text_frame
    tf.margin_left = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "IMPLEMENTED HOST-LEVEL MITIGATION"
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    mech_items = [
        ("Cross-Platform Command Dispatch", "Implemented in `backend/services/response_executor.py`:\n• Linux: `sudo iptables -A INPUT -s <IP> -j DROP`\n• Windows: `netsh advfirewall firewall add rule name=\"PhantomNet_AutoBlock_<IP>\" dir=in action=block remoteip=<IP>`"),
        ("Dual-Threshold Trigger Gate", "Mitigation actions only execute when threat score exceeds the strict `0.8` confidence gate (Ensemble Block Mode with 0.0% false positives)."),
        ("Sliding-Window Rate Limiting", "Maintains an in-memory window of request timestamps to enforce maximum requests-per-minute per remote IP address."),
        ("Audit Logging & Tracking", "Every mitigation attempt creates a structured audit record containing timestamp, IP, trigger threat level, and platform.")
    ]
    for title, desc in mech_items:
        p_t = tf.add_paragraph()
        p_t.text = f"\n• {title}"
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = BLUE_ACCENT
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MAIN

    right_c = add_card(slide10, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf_r = right_c.text_frame
    tf_r.margin_left = tf_r.margin_top = Inches(0.25)
    pr = tf_r.paragraphs[0]
    pr.text = "ENGINEERING BOUNDARIES & AUDIT FINDINGS"
    pr.font.name = "Calibri"
    pr.font.size = Pt(14)
    pr.font.bold = True
    pr.font.color.rgb = AMBER_WARN

    bound_items = [
        ("Volatile In-Memory Tracking", "Blocked IP records are stored in an in-memory dictionary (`self.blocked_ips`). Blocks are not persisted across backend restarts in PostgreSQL or Redis."),
        ("Optimistic Execution Status", "The method returns `status: blocked` based on command dispatch; it does not verify whether the container had sufficient Linux capabilities (`NET_ADMIN`) to execute iptables."),
        ("Honeypot Decoupling", "Honeypot services do not check the application blocklist; traffic suppression relies entirely on whether the host-level OS firewall dropped the packets."),
        ("Stubbed Features (Future Work)", "• Automatic honeypot container scaling (`scale_honeypots`) is currently a logging stub.\n• Administrator paging (`notify_admin`) is currently a logging stub.\n• Distributed cluster-wide block propagation is planned.")
    ]
    for title, desc in bound_items:
        p_t = tf_r.add_paragraph()
        p_t.text = f"\n• {title}"
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = RED_ALERT
        p_d = tf_r.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MUTED

    add_notes(slide10, "Examiner Key Point:\n'When asked: Does PhantomNet automatically scale containers or enforce Kubernetes network policies? Answer: No. We do not claim Kubernetes orchestration. The active defense layer issues host-level iptables and netsh commands, and container auto-scaling is clearly scoped as future engineering work.'")

    # =========================================================================
    # SLIDE 11: Sentinel Automated Playbook Generation
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, BG_LIGHT)
    add_header(slide11, "Sentinel — Automated Detection Engineering", "Incident Response", "11 / 18")

    sentinel_cards = [
        ("🎯 MITRE ATT&CK Mapping", "Enterprise TTP Correlation", "Automated correlation of honeypot telemetry with official MITRE ATT&CK techniques:\n• T1046: Network Service Discovery\n• T1110.001: Password Guessing / Brute Force\n• T1190: Exploit Public-Facing Application\n• T1595: Active Scanning & Reconnaissance\nPopulates MITRE Navigator compatible matrices."),
        ("📐 Sigma Detection Rules", "Vendor-Agnostic SIEM Rules", "Generates compliant Sigma YAML detection rules per incident:\n• Standardized logsource mappings\n• Specific attacker IOC selectors (IP, port, headers)\n• Severity-ranked condition logic\n• Import-ready for Splunk, Elastic, QRadar, and Microsoft Sentinel."),
        ("📋 Snort Network Rules", "Network IDS Signatures", "Synthesizes import-ready Snort IDS rules:\n• Protocols: TCP, HTTP, SMTP, FTP\n• Auto-incrementing unique SIDs\n• Directional flow signatures (`$EXTERNAL_NET` to `$HOME_NET`)\n• Reference tags linking to ATT&CK technique IDs."),
        ("📦 STIX 2.1 Threat Intel", "TAXII 2.1 Sharing Feeds", "Produces fully structured STIX 2.1 JSON bundles:\n• Indicator, Attack Pattern, and Identity SDOs\n• Structured relationship objects\n• Integration with TAXII 2.1 collections\n• Exportable to MISP and AlienVault OTX.")
    ]
    for i, (title, sub, body) in enumerate(sentinel_cards):
        x = Inches(0.8 + (i % 2) * 5.95)
        y = Inches(1.5 + (i // 2) * 2.7)
        c = add_card(slide11, x, y, Inches(5.75), Inches(2.55))
        tf = c.text_frame
        tf.margin_left = tf.margin_top = Inches(0.2)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Calibri"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = NAVY_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = CYAN_TECH
        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.name = "Calibri"
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = TEXT_MAIN

    add_notes(slide11, "Examiner Key Point:\n'Sentinel is one of the standout accomplishments of PhantomNet. Rather than just alerting that an attack occurred, it programmatically converts the raw attack telemetry into formal detection rules (Sigma/Snort) and threat-sharing bundles (STIX 2.1) that SOC analysts can immediately push into production firewalls and SIEMs.'")

    # =========================================================================
    # SLIDE 12: LLM Graceful Degradation Architecture
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, BG_LIGHT)
    add_header(slide12, "LLM Layer — Graceful Degradation Architecture", "System Resilience", "12 / 18")

    top_b = add_card(slide12, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.75), bg_color=TAG_BG_BLUE, border_color=CYAN_TECH)
    tf = top_b.text_frame
    tf.margin_left = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "CORE DEFENSIVE PRINCIPLE: LLMs provide narrative enrichment — NOT mission-critical detection logic. The incident response pipeline remains 100% operational whether the local LLM is available or completely offline."
    p.font.name = "Calibri"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    path_a = add_card(slide12, Inches(0.8), Inches(2.45), Inches(5.7), Inches(3.6), bg_color=CARD_BG, border_color=BLUE_ACCENT)
    tf_a = path_a.text_frame
    tf_a.margin_left = tf_a.margin_top = Inches(0.25)
    pa = tf_a.paragraphs[0]
    pa.text = "PATH A: OLLAMA LLM AVAILABLE"
    pa.font.name = "Calibri"
    pa.font.size = Pt(13)
    pa.font.bold = True
    pa.font.color.rgb = BLUE_ACCENT

    items_a = [
        "Inference Engine: Local Ollama container running `mistral` 7B.",
        "Strict Timeouts: Configured with async timeouts to prevent API blocking.",
        "Narrative Synthesis: Enriches playbooks with human-readable threat actor summaries and contextual remediation advice.",
        "Redis Caching: Generated narratives cached for 24 hours (86,400s TTL) by threat context hash to avoid redundant GPU inference."
    ]
    for it in items_a:
        p_it = tf_a.add_paragraph()
        p_it.text = f"• {it}"
        p_it.font.name = "Calibri"
        p_it.font.size = Pt(11)
        p_it.font.color.rgb = TEXT_MAIN

    path_b = add_card(slide12, Inches(6.8), Inches(2.45), Inches(5.7), Inches(3.6), bg_color=CARD_BG, border_color=GREEN_SUCCESS)
    tf_b = path_b.text_frame
    tf_b.margin_left = tf_b.margin_top = Inches(0.25)
    pb = tf_b.paragraphs[0]
    pb.text = "PATH B: DETERMINISTIC JINJA2 FALLBACK"
    pb.font.name = "Calibri"
    pb.font.size = Pt(13)
    pb.font.bold = True
    pb.font.color.rgb = GREEN_SUCCESS

    items_b = [
        "Trigger Conditions: `SENTINEL_LLM_ENABLED=false`, Ollama container offline, or HTTP request timeout.",
        "Deterministic Engine: `generate_fallback_narrative()` compiles structured Markdown via Jinja2 templates.",
        "Zero Downtime: Pipeline proceeds immediately with zero exception leakage and sub-millisecond execution.",
        "Identical Security Artifacts: Sigma rules, Snort rules, and STIX 2.1 bundles generate identically across both paths."
    ]
    for it in items_b:
        p_it = tf_b.add_paragraph()
        p_it.text = f"• {it}"
        p_it.font.name = "Calibri"
        p_it.font.size = Pt(11)
        p_it.font.color.rgb = TEXT_MAIN

    unify_box = add_card(slide12, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.8), bg_color=TAG_BG_GREEN, border_color=GREEN_SUCCESS)
    tf_u = unify_box.text_frame
    tf_u.margin_left = Inches(0.25)
    tf_u.vertical_anchor = MSO_ANCHOR.MIDDLE
    pu = tf_u.paragraphs[0]
    pu.text = "VERIFIED IN CODEBASE (`backend/sentinel/llm_service.py`): Both paths feed into the same database models (`SentinelPlaybook`), API routes, and frontend views. The system exhibits zero hard dependency on LLM hardware."
    pu.font.name = "Calibri"
    pu.font.size = Pt(11)
    pu.font.bold = True
    pu.font.color.rgb = GREEN_SUCCESS

    add_notes(slide12, "Examiner Key Point:\n'Many student AI projects break completely if their LLM server crashes. In PhantomNet, we implemented explicit graceful degradation. If Ollama times out or is disabled, our Jinja2 fallback produces complete, valid playbooks immediately. The defensive capability is independent of the generative AI component.'")

    # =========================================================================
    # SLIDE 13: SOC Observability Dashboard — Verified Frontend
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13, BG_LIGHT)
    add_header(slide13, "SOC Observability Dashboard — Verified Frontend", "User Interface", "13 / 18")

    c_left = add_card(slide13, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf = c_left.text_frame
    tf.margin_left = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "REACT 19 SINGLE PAGE APPLICATION"
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p_build = tf.add_paragraph()
    p_build.text = "• Stack: React 19 • Vite (Rolldown) • Tailwind CSS • Recharts • Leaflet • @xyflow/react\n• Production Build: Built in 11.65 seconds (3,035 modules transformed to `dist/`)\n"
    p_build.font.name = "Calibri"
    p_build.font.size = Pt(10.5)
    p_build.font.color.rgb = CYAN_TECH

    p_views_head = tf.add_paragraph()
    p_views_head.text = "11 VERIFIED PRODUCTION VIEWS:"
    p_views_head.font.name = "Calibri"
    p_views_head.font.size = Pt(11)
    p_views_head.font.bold = True
    p_views_head.font.color.rgb = NAVY_PRIMARY

    views_list = [
        "1. Dashboard.jsx — Real-time event rate gauge & threat summary",
        "2. Events.jsx — Paginated, searchable interaction event log",
        "3. Honeypots.jsx — Protocol status, port bindings, and capture activity",
        "4. ThreatAnalysis.jsx — Threat scores, feature vectors, SHAP values",
        "5. SentinelDashboard.jsx — Playbooks, Sigma/Snort rules, STIX export",
        "6. ThreatHunting.jsx — Ad-hoc IOC search, correlation filters, timelines",
        "7. AdminPanel.jsx — User management, RBAC configuration, audit logs",
        "8. AdvancedAnalytics.jsx — Longitudinal trend analysis & graphs",
        "9. AnomalyDashboard.jsx — Statistical anomaly and cluster viewer",
        "10. Topology.jsx — Interactive network node graph (@xyflow/react)",
        "11. Login.jsx — JWT-authenticated secure access control"
    ]
    for v in views_list:
        pv = tf.add_paragraph()
        pv.text = v
        pv.font.name = "Calibri"
        pv.font.size = Pt(9.5)
        pv.font.color.rgb = TEXT_MAIN

    screenshot_path = "docs/screenshots/dashboard.png"
    if os.path.exists(screenshot_path):
        c_right = add_card(slide13, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
        slide13.shapes.add_picture(screenshot_path, Inches(6.9), Inches(1.65), Inches(5.5), Inches(4.3))
        cap_box = slide13.shapes.add_textbox(Inches(6.8), Inches(6.05), Inches(5.7), Inches(0.75))
        tf_c = cap_box.text_frame
        tf_c.word_wrap = True
        pc = tf_c.paragraphs[0]
        pc.text = "Verified Production Dashboard UI: Live event telemetry, real-time threat gauges, and protocol activity monitors connected via WebSocket (`/api/v1/realtime/ws`)."
        pc.font.name = "Calibri"
        pc.font.size = Pt(10)
        pc.font.color.rgb = TEXT_MUTED
    else:
        c_right = add_card(slide13, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
        tf_r = c_right.text_frame
        tf_r.margin_left = tf_r.margin_top = Inches(0.25)
        pr = tf_r.paragraphs[0]
        pr.text = "ANALYST CONTROL & WORKFLOW CAPABILITIES"
        pr.font.name = "Calibri"
        pr.font.size = Pt(13)
        pr.font.bold = True
        pr.font.color.rgb = NAVY_PRIMARY

    add_notes(slide13, "Examiner Key Point:\n'Correction note: The old presentation claimed 14 pages. Our codebase audit verified exactly 11 distinct functional views in frontend-dev/phantomnet-dashboard/src/pages. The application was compiled to production in 11.65 seconds with zero build warnings, verifying that the frontend is fully implemented.'")

    # =========================================================================
    # SLIDE 14: SDN Emulation & Deployment Architecture
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14, BG_LIGHT)
    add_header(slide14, "SDN Emulation & Deployment Architecture", "Deployment & Network", "14 / 18")

    left_c = add_card(slide14, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf = left_c.text_frame
    tf.margin_left = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "SDN CONTROLLER & TOPOLOGY EMULATION"
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    sdn_items = [
        ("POX OpenFlow Controller Extension", "Implemented in `pox/ext/phantomnet_controller.py` (661 lines of Python):\n• OpenFlow 1.0 multi-switch flow rule management\n• Protocol-based round-robin honeypot pools (SSH, HTTP, SMTP, FTP)\n• Session persistence via MAC/IP binding tables\n• Dynamic redirection and health monitoring"),
        ("14-Host Mininet Emulation Topology", "Defined in `topology/phantomnet_topology.py`:\n• 14 hosts across 2 Open vSwitch switches (s1, s2)\n• Protocol pools: SSH (h1, h7, h10), HTTP (h2, h8, h11)\n• Dedicated coordinator node (h5) & attacker node (h6)"),
        ("Operating Environment Boundary", "Mininet and POX require Linux kernel network namespaces and Open vSwitch. On Windows host machines, attack scenarios run through the socket-based local simulator (`simulation/attack_campaign_local.py`).")
    ]
    for title, desc in sdn_items:
        pt = tf.add_paragraph()
        pt.text = f"\n• {title}"
        pt.font.name = "Calibri"
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = BLUE_ACCENT
        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Calibri"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_MAIN

    right_c = add_card(slide14, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf_r = right_c.text_frame
    tf_r.margin_left = tf_r.margin_top = Inches(0.25)
    pr = tf_r.paragraphs[0]
    pr.text = "DOCKER COMPOSE ORCHESTRATION"
    pr.font.name = "Calibri"
    pr.font.size = Pt(13)
    pr.font.bold = True
    pr.font.color.rgb = NAVY_PRIMARY

    services_info = [
        ("Database & Broker", "• `postgres` (15-alpine): WAL archive, healthcheck\n• `redis` (7-alpine): AOF persistence, noeviction"),
        ("Honeypot Microservices", "• `ssh_honeypot`: Port 2722:2222\n• `http_honeypot`: Port 8080:8080\n• `ftp_honeypot`: Port 2721:2121\n• `smtp_honeypot`: Port 2725:2525"),
        ("Backend & LLM", "• `api`: FastAPI (8000), NET_RAW capability\n• `ollama`: Mistral 7B inference service (11434)"),
        ("Frontend Delivery", "• `frontend`: Nginx serving production Vite build (3000:8080)"),
        ("Isolated Network Bridges", "• `honeypot_net` (internal: true)\n• `internal_broker_net` (internal: true)\n• `app_net` (bridged frontend/API communication)")
    ]
    for title, desc in services_info:
        pt = tf_r.add_paragraph()
        pt.text = f"\n• {title}"
        pt.font.name = "Calibri"
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = CYAN_TECH
        pd = tf_r.add_paragraph()
        pd.text = desc
        pd.font.name = "Calibri"
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MAIN

    add_notes(slide14, "Examiner Key Point:\n'Be transparent about SDN: Do NOT claim live hardware SDN switches. Explain that POX and Mininet were implemented in 661 lines of controller code as a Linux network emulation testbed. On Windows, socket-level attack campaigns test the same port mappings.'")

    # =========================================================================
    # SLIDE 15: Security & Container Hardening Controls
    # =========================================================================
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15, BG_LIGHT)
    add_header(slide15, "Security & Container Hardening Controls", "Defensive Controls", "15 / 18")

    sec_cards = [
        ("🔒 Container Hardening", "Least-Privilege Isolation", "Verified in `docker-compose.yml`:\n• `cap_drop: ALL` drops all Linux root capabilities\n• `security_opt: no-new-privileges:true` prevents escalation\n• `read_only: true` locks root filesystem against tampering\n• `tmpfs` mounts (`/tmp`, `/run`) for non-persistent execution\n• Non-root UID/GID (`10001:10001`) execution.", CYAN_TECH),
        ("🌐 Network Segmentation", "Internal Bridge Isolation", "• Honeypots partitioned in `honeypot_net` (`internal: true`)\n• Zero external egress from honeypot containers\n• Zero direct access to PostgreSQL database or Redis\n• Honeypots communicate with API strictly via HMAC-signed POST\n• Backend services partitioned on `internal_broker_net`.", BLUE_ACCENT),
        ("🔑 Access Control & Auth", "RBAC & Cryptographic Signing", "• JWT Bearer tokens enforced across API routes\n• Role-Based Access Control: Admin, Analyst, Viewer\n• Pydantic v2 strict schema validation on all payloads\n• CSRF tokens validated on all state-modifying requests\n• Rate limiting applied on sensitive management endpoints.", GREEN_SUCCESS),
        ("🛡️ Application Defenses", "Security Headers & Sanitization", "• Security headers: CSP, X-Frame-Options, X-XSS-Protection\n• Credential Sanitizer (`credential_sanitizer.py`) masks sensitive plaintexts\n• Local Event Spooler protects against buffer exhaustion\n• Comprehensive audit trail for all admin operations.", NAVY_PRIMARY)
    ]
    for i, (title, sub, body, color) in enumerate(sec_cards):
        x = Inches(0.8 + (i % 2) * 5.95)
        y = Inches(1.5 + (i // 2) * 2.7)
        c = add_card(slide15, x, y, Inches(5.75), Inches(2.55))
        tf = c.text_frame
        tf.margin_left = tf.margin_top = Inches(0.2)
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Calibri"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = NAVY_PRIMARY
        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.name = "Calibri"
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = TEXT_MAIN

    add_notes(slide15, "Examiner Key Point:\n'Emphasize container containment: If an attacker drops an exploit payload inside the HTTP or SSH honeypot, the container root filesystem is read-only, all Linux capabilities are dropped, and the internal Docker bridge prevents the container from routing packets to the Internet or the database.'")

    # =========================================================================
    # SLIDE 16: Comprehensive Testing & Empirical Validation
    # =========================================================================
    slide16 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide16, BG_LIGHT)
    add_header(slide16, "Testing & Quality Assurance — Empirical Findings", "Verification & QA", "16 / 18")

    left_c = add_card(slide16, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf = left_c.text_frame
    tf.margin_left = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "FULL SUITE MONOLITHIC EXECUTION"
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p_stats = tf.add_paragraph()
    p_stats.text = "• 3,848 Passed\n• 166 Failed\n• 232 Errors\n• 2 Skipped (63,800+ lines of test output)\n"
    p_stats.font.name = "Calibri"
    p_stats.font.size = Pt(11)
    p_stats.font.bold = True
    p_stats.font.color.rgb = AMBER_WARN

    p_rc = tf.add_paragraph()
    p_rc.text = "ROOT CAUSE OF MONOLITHIC TEST ERRORS:"
    p_rc.font.name = "Calibri"
    p_rc.font.size = Pt(11)
    p_rc.font.bold = True
    p_rc.font.color.rgb = NAVY_PRIMARY

    errors_breakdown = [
        "1. SQLite Index Collisions (144 errors):\nShared in-memory SQLite instances in multi-file suite runs collided on `ix_sentinel_playbooks_src_ip already exists`.",
        "2. SQLAlchemy Mapper Collisions (65 errors):\nCross-module model re-declarations caused mapper conflicts when running all 176 test files in one pytest process.",
        "3. Missing CI Security Gate File (Fixture failure):\n`.github/workflows/security_quality_gates.yml` assertion failed in `test_ci_security_gates.py`."
    ]
    for eb in errors_breakdown:
        pe = tf.add_paragraph()
        pe.text = eb
        pe.font.name = "Calibri"
        pe.font.size = Pt(10)
        pe.font.color.rgb = TEXT_MUTED

    right_c = add_card(slide16, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf_r = right_c.text_frame
    tf_r.margin_left = tf_r.margin_top = Inches(0.25)
    pr = tf_r.paragraphs[0]
    pr.text = "ISOLATED UNIT TEST SUITE EXECUTION"
    pr.font.name = "Calibri"
    pr.font.size = Pt(13)
    pr.font.bold = True
    pr.font.color.rgb = GREEN_SUCCESS

    pr_stats = tf_r.add_paragraph()
    pr_stats.text = "• 931 PASSED\n• 1 FAILED (Execution time: 9.34 seconds)\n"
    pr_stats.font.name = "Calibri"
    pr_stats.font.size = Pt(12)
    pr_stats.font.bold = True
    pr_stats.font.color.rgb = GREEN_SUCCESS

    pr_detail = tf_r.add_paragraph()
    pr_detail.text = "EVALUATION SIGNIFICANCE:\n• Demonstrates that core business logic, feature engineering, and playbook synthesis execute with near-perfect reliability when isolated from monolithic database fixture cross-talk.\n• Single failure in `tests/unit`: Minor assertion mismatch on default deception label ('balanced_deception' vs 'aggressive_deception').\n\nENGINEERING LESSON & REFACTORING SCOPE:\n• Test isolation: Replace global shared SQLite fixtures with per-test ephemeral database containers or transaction-rollback fixtures.\n• Parallel execution: Enforce pytest-xdist safe fixture isolation."
    pr_detail.font.name = "Calibri"
    pr_detail.font.size = Pt(10.5)
    pr_detail.font.color.rgb = TEXT_MAIN

    add_notes(slide16, "Examiner Key Point:\n'Academic honesty in testing: We do not hide our 232 test errors. We transparently explain that 144 errors were caused by SQLite index re-declaration when all 176 test files ran in a single process. When isolated by test domain (such as pytest tests/unit), 931 out of 932 tests pass cleanly in 9.3 seconds, confirming application correctness.'")

    # =========================================================================
    # SLIDE 17: Current Limitations & Future Scope
    # =========================================================================
    slide17 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide17, BG_LIGHT)
    add_header(slide17, "Current Limitations & Future Engineering Scope", "Engineering Boundaries", "17 / 18")

    c_left = add_card(slide17, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf = c_left.text_frame
    tf.margin_left = tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "CURRENT LIMITATIONS & BOUNDARIES"
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = AMBER_WARN

    limits = [
        ("1. ML Domain Specificity", "Models are optimized for the synthetic/honeypot feature space. Zero-shot transfer to heterogeneous enterprise datasets collapses without domain adaptation."),
        ("2. Volatile Active Defense State", "Blocked IP lists reside in memory (`self.blocked_ips`); persistent database storage in PostgreSQL/Redis is not yet implemented."),
        ("3. Optimistic Firewall Return", "Firewall block status is reported optimistically upon subprocess dispatch without container-level privilege verification."),
        ("4. Stubbed Elastic Scaling", "Honeypot container scaling (`scale_honeypots`) and administrator notifications are currently logging stubs."),
        ("5. Monolithic Test Isolation", "Test fixtures require database isolation refactoring to prevent SQLite index collision during full-suite runs.")
    ]
    for title, desc in limits:
        pt = tf.add_paragraph()
        pt.text = f"\n• {title}"
        pt.font.name = "Calibri"
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = RED_ALERT
        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Calibri"
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED

    c_right = add_card(slide17, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    tf_r = c_right.text_frame
    tf_r.margin_left = tf_r.margin_top = Inches(0.25)
    pr = tf_r.paragraphs[0]
    pr.text = "FUTURE ENGINEERING & RESEARCH SCOPE"
    pr.font.name = "Calibri"
    pr.font.size = Pt(13)
    pr.font.bold = True
    pr.font.color.rgb = GREEN_SUCCESS

    futures = [
        ("1. Transfer Learning & Feature Alignment", "Implement domain adaptation algorithms (CORAL, adversarial domain adaptation) and retrain across heterogeneous public datasets (CIC-IDS, UNSW)."),
        ("2. Persistent Distributed Blocklists", "Migrate blocklists to PostgreSQL with Redis caching; distribute block rules cluster-wide via Redis Pub/Sub."),
        ("3. Verified Firewall Feedback Loops", "Verify firewall return codes and add netfilter kernel inspection before confirming active mitigation."),
        ("4. Dynamic Container Elasticity", "Implement Docker SDK / Kubernetes API lifecycle management to spin up dynamic honeypot replicas on demand."),
        ("5. Isolated Parallel Test Infrastructure", "Refactor test fixtures with transactional rollbacks and isolated PostgreSQL test instances.")
    ]
    for title, desc in futures:
        pt = tf_r.add_paragraph()
        pt.text = f"\n• {title}"
        pt.font.name = "Calibri"
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = BLUE_ACCENT
        pd = tf_r.add_paragraph()
        pd.text = desc
        pd.font.name = "Calibri"
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MAIN

    add_notes(slide17, "Examiner Key Point:\n'A strong project does not claim to have solved every problem. A strong project clearly demarcates what was implemented, what was measured, what failed, and what the precise engineering steps are to solve those limitations.'")

    # =========================================================================
    # SLIDE 18: Project Summary & Measurable Outcomes
    # =========================================================================
    slide18 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide18, NAVY_DARK)

    cat_box = slide18.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.0), Inches(0.35))
    tf_cat = cat_box.text_frame
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = "CONCLUSION & SUMMARY"
    p_cat.font.name = "Calibri"
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = CYAN_TECH

    title_box = slide18.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.65))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = "What PhantomNet Demonstrates"
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    achievements = [
        ("4", "Custom Micro-Honeypots", "SSH, HTTP, FTP, SMTP with HMAC-signed telemetry and spooling"),
        ("160", "API Route Endpoints", "Across 27 FastAPI routers with 2 real-time WebSocket feeds"),
        ("97.88%", "In-Domain RF Accuracy", "F1: 96.41%, AUC-ROC: 0.9968 on in-distribution synthetic flows"),
        ("0.0%", "False Positive Rate", "100.0% Precision at 0.8 safety blocking gate on held-out test data"),
        ("ATT&CK + Sigma", "Detection Engineering", "Synthesizes Sigma YAML, Snort IDS rules, and STIX 2.1 bundles"),
        ("React 19", "Production SOC UI", "11 verified views, built in 11.65 seconds (3,035 modules)")
    ]
    for i, (bignum, label, subtext) in enumerate(achievements):
        row = i // 3
        col = i % 3
        x = Inches(0.8 + col * 3.95)
        y = Inches(1.65 + row * 2.2)
        c = add_card(slide18, x, y, Inches(3.75), Inches(1.95), bg_color=RGBColor(0x16, 0x24, 0x3C), border_color=RGBColor(0x2A, 0x43, 0x68))
        tf = c.text_frame
        tf.margin_left = tf.margin_top = Inches(0.2)
        p1 = tf.paragraphs[0]
        p1.text = bignum
        p1.font.name = "Calibri"
        p1.font.size = Pt(32)
        p1.font.bold = True
        p1.font.color.rgb = CYAN_TECH
        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.name = "Calibri"
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p3 = tf.add_paragraph()
        p3.text = subtext
        p3.font.name = "Calibri"
        p3.font.size = Pt(10)
        p3.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    thesis_box = add_card(slide18, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9), bg_color=RGBColor(0x1E, 0x2E, 0x48), border_color=CYAN_TECH)
    tf_th = thesis_box.text_frame
    tf_th.margin_left = Inches(0.25)
    tf_th.vertical_anchor = MSO_ANCHOR.MIDDLE
    pth = tf_th.paragraphs[0]
    pth.text = "FINAL THESIS STATEMENT:\n'PhantomNet demonstrates that deception, telemetry, and machine learning can be combined into a safe, explainable, and resilient threat intelligence architecture from initial adversary interaction to analyst-ready intelligence.'"
    pth.font.name = "Calibri"
    pth.font.size = Pt(11.5)
    pth.font.bold = True
    pth.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    add_notes(slide18, "Final Closing Statement:\n'To conclude: PhantomNet is not a startup pitch claiming to solve all of cybersecurity. It is an honest, audited, and empirically evaluated research platform that integrates four custom deception protocols, 160 API endpoints, a zero-false-positive safety gate for automated mitigation, resilient dual-mode playbook synthesis, and a production React 19 SOC dashboard. Thank you, and we welcome your questions.'")

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_file = "PhantomNet_Final_Presentation.pptx"
    if len(sys.argv) > 1:
        out_file = sys.argv[1]
    build_presentation(out_file)
